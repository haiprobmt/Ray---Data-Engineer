"""Bounded document text readers. Content is evidence, never executable instructions."""
from __future__ import annotations

import hashlib
import io
import json
import logging
from pathlib import Path
import subprocess
import sys
import time
import zipfile

from .memory import redact

MAX_BYTES = 20_000_000
MAX_TEXT = 80_000
EXTENSIONS = {".md", ".txt", ".csv", ".docx", ".pdf", ".xlsx", ".xls"}


class ReadError(ValueError):
    """Only fixed, credential-free explanations cross the reader boundary."""


def check_file(name, size):
    if not isinstance(name, str) or Path(name).suffix.lower() not in EXTENSIONS:
        raise ReadError("Send an MD, DOCX, PDF, XLSX, XLS, CSV or TXT file.")
    if type(size) is not int or not 0 <= size <= MAX_BYTES:
        raise ReadError("This file is too large. Send a file smaller than 20 MB.")


def checked_zip(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if (len(entries) > 5000 or sum(e.file_size for e in entries) > 64_000_000
            or any(e.file_size > 1_000_000 and e.file_size > max(e.compress_size, 1) * 200 for e in entries)):
        archive.close()
        raise ReadError("This file expands beyond the reading limit. Split it into smaller files.")
    return archive


def extract_document(name, data):
    """Parser worker implementation. Call read_document at untrusted intake boundaries."""
    check_file(name, len(data))
    result = {"source": "document", "name": redact(Path(name.replace("\\", "/")).name)[:180],
              "sha256": hashlib.sha256(data).hexdigest(), "sections": [], "warnings": [], "truncated": False}
    used = 0

    def add(location, text):
        nonlocal used
        text = redact(str(text))
        location = redact(str(location))[:240]
        if not text.strip():
            return True
        remaining = MAX_TEXT - used - len(location)
        if remaining <= 0 or len(result["sections"]) >= 1000:
            result["truncated"] = True
            return False
        result["sections"].append({"location": location, "text": text[:remaining]})
        used += len(location) + min(len(text), remaining)
        if len(text) > remaining:
            result["truncated"] = True
            return False
        return True

    try:
        suffix = Path(name).suffix.lower()
        if suffix in {".md", ".txt", ".csv"}:
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            add("Text", data.decode(encoding))
        elif suffix == ".docx":
            from defusedxml import ElementTree
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            result["warnings"].append("Text and tables only; pictures, text boxes and layout are not inspected.")
            with checked_zip(data) as archive:
                parts = ["word/document.xml"] + sorted(n for n in archive.namelist() if n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and n.endswith(".xml"))
                for part in parts:
                    root = ElementTree.fromstring(archive.read(part))
                    body = root.find("w:body", ns) if part == "word/document.xml" else root
                    if body is None:
                        raise ReadError("The Word document has no readable body.")
                    for number, block in enumerate(body, 1):
                        kind = "table" if block.tag.endswith("}tbl") else "paragraph"
                        # Newlines preserve paragraph and cell boundaries, including inside tables.
                        paragraphs = [block] if block.tag.endswith("}p") else block.findall(".//w:p", ns)
                        text = "\n".join("".join(p.itertext()) for p in paragraphs)
                        if not add(f"{part}: {kind} {number}", text):
                            break
                    if result["truncated"]:
                        break
        elif suffix == ".pdf":
            from pypdf import PdfReader
            logging.getLogger("pypdf").setLevel(logging.CRITICAL)
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ReadError("This PDF is password protected. Send an unlocked copy without sharing the password.")
            result["page_count"] = len(reader.pages)
            result["warnings"].append("PDF text only; pictures and visual layout are not inspected.")
            empty = []
            for number, page in enumerate(reader.pages[:200], 1):
                text = page.extract_text() or ""
                if not text.strip():
                    empty.append(number)
                if not add(f"Page {number}", text):
                    break
            if empty:
                result["warnings"].append("Some pages have no readable text. Scanned pages need OCR; blank pages may also have no text.")
            result["truncated"] |= len(reader.pages) > 200
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            with checked_zip(data):
                pass
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            result["warnings"].append("Cell values and formulas only; formulas are not calculated, and charts or pictures are not inspected.")
            try:
                for sheet in wb.worksheets[:40]:
                    if sheet.sheet_state != "visible":
                        result["warnings"].append(f"Hidden sheet included: {redact(sheet.title)}")
                    if sheet.max_row is None or sheet.max_column is None:
                        sheet.calculate_dimension(force=True)
                    result["truncated"] |= sheet.max_row > 5000 or sheet.max_column > 100
                    for row in sheet.iter_rows(max_row=min(sheet.max_row, 5000), max_col=min(sheet.max_column, 100)):
                        cells = [f"{c.coordinate}={c.value}" for c in row if c.value is not None]
                        if cells and not add(f"Sheet {sheet.title}", " | ".join(cells)):
                            break
                    if used >= MAX_TEXT:
                        break
                result["truncated"] |= len(wb.worksheets) > 40
            finally:
                wb.close()
        elif suffix == ".xls":
            import xlrd
            wb = xlrd.open_workbook(file_contents=data, on_demand=True)
            result["warnings"].append("Legacy Excel: saved cell values only; formulas, charts and pictures are not inspected.")
            try:
                from openpyxl.utils import get_column_letter
                for sheet in wb.sheets()[:40]:
                    if sheet.visibility:
                        result["warnings"].append(f"Hidden sheet included: {redact(sheet.name)}")
                    result["truncated"] |= sheet.nrows > 5000 or sheet.ncols > 100
                    for row in range(min(sheet.nrows, 5000)):
                        cells = []
                        for col in range(min(sheet.ncols, 100)):
                            cell = sheet.cell(row, col)
                            value = xlrd.xldate_as_datetime(cell.value, wb.datemode).isoformat() if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                            if value != "":
                                cells.append(f"{get_column_letter(col+1)}{row+1}={value}")
                        if cells and not add(f"Sheet {sheet.name}", " | ".join(cells)):
                            break
                    if used >= MAX_TEXT:
                        break
                result["truncated"] |= wb.nsheets > 40
            finally:
                wb.release_resources()
    except ReadError:
        raise
    except Exception:
        raise ReadError("I couldn't read this file. It may be damaged, locked, or saved in a different format. Try exporting a fresh copy.") from None
    if result["truncated"]:
        result["warnings"].append("Only part of this file was read. Split the file or send the relevant pages or sheets for a complete check.")
    if not result["sections"]:
        result["warnings"].append("No readable text or cell values were found.")
    return result


def read_document(name, data, *, cancel=None):
    """Parse in a short-lived process so malformed files cannot hang the gateway."""
    from .runtime import clean_env
    check_file(name, len(data))
    if cancel:
        cancel.check()
    process = subprocess.Popen([sys.executable, "-B", "-m", "ray_de.documents", "document" + Path(name).suffix.lower()],
                               stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                               cwd=Path(__file__).resolve().parents[1],
                               env=clean_env(), creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    started = time.monotonic()
    payload = data
    try:
        while True:
            if cancel:
                cancel.check()
            if time.monotonic() - started > 30:
                raise ReadError("Reading this file took too long. Send a smaller file or export a fresh copy.")
            try:
                output, _ = process.communicate(payload, timeout=0.2)
                break
            except subprocess.TimeoutExpired:
                payload = None
        try:
            result = json.loads(output)
            if result.get("error"):
                raise ReadError(result["error"])
            if process.returncode or not isinstance(result.get("sections"), list):
                raise ValueError()
            result["name"] = redact(Path(name.replace("\\", "/")).name)[:180]
            return result
        except ReadError:
            raise
        except Exception:
            raise ReadError("I couldn't finish reading this file. Try a smaller file or a fresh export.") from None
    finally:
        if process.poll() is None:
            process.kill()
        process.communicate()


if __name__ == "__main__":
    try:
        output = extract_document(sys.argv[1], sys.stdin.buffer.read(MAX_BYTES + 1))
    except ReadError as exc:
        output = {"error": str(exc)}
    print(json.dumps(output, ensure_ascii=True))
